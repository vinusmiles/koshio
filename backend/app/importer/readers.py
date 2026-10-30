"""Turn an uploaded file into a list of rows of text cells, entirely in memory."""
import csv
import io
from datetime import date, datetime, time

MAX_BYTES = 5 * 1024 * 1024
MAX_PDF_BYTES = 15 * 1024 * 1024  # card statements carry large promotional images
MAX_ROWS = 20_000
SUPPORTED = (".csv", ".txt", ".xlsx", ".xls", ".pdf")
LATER = {".ofx": "OFX files", ".qfx": "QFX files"}


class ReadError(ValueError):
    """A problem with the file itself, shown to the user as is."""


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d %H:%M:%S") if v.time() != time(0) else v.strftime("%Y-%m-%d")
    if isinstance(v, date):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float):
        return f"{v:.2f}" if v != int(v) else str(int(v))
    return str(v).strip()


def _trim(rows: list[list[str]]) -> list[list[str]]:
    rows = [[c.strip() for c in r] for r in rows]
    if len(rows) > MAX_ROWS:
        raise ReadError(f"This file has more than {MAX_ROWS:,} rows. Split it into smaller statements.")
    return rows


def _csv(data: bytes) -> list[list[str]]:
    for enc in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    sample = text[:20000]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return _trim(list(csv.reader(io.StringIO(text), dialect)))


def _xlsx(data: bytes) -> list[list[str]]:
    from openpyxl import load_workbook
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ReadError("This Excel file couldn't be opened. It may be damaged or password-protected.")
    ws = _first_non_empty(wb.worksheets, lambda s: s.max_row and s.max_row > 1) or wb.active
    rows = [[_cell(v) for v in r] for r in ws.iter_rows(values_only=True)]
    wb.close()
    return _trim(rows)


def _xls(data: bytes) -> list[list[str]]:
    import xlrd
    try:
        book = xlrd.open_workbook(file_contents=data)
    except Exception:
        raise ReadError("This .xls file couldn't be opened. Try saving it as .xlsx or .csv.")
    sheet = _first_non_empty(book.sheets(), lambda s: s.nrows > 1) or book.sheet_by_index(0)
    rows = []
    for r in range(sheet.nrows):
        out = []
        for c in sheet.row(r):
            if c.ctype == xlrd.XL_CELL_DATE:
                out.append(_cell(xlrd.xldate.xldate_as_datetime(c.value, book.datemode)))
            else:
                out.append(_cell(c.value))
        rows.append(out)
    return _trim(rows)


def _first_non_empty(items, test):
    return next((i for i in items if test(i)), None)


def read_rows(filename: str, data: bytes) -> list[list[str]]:
    name = (filename or "").lower()
    ext = name[name.rfind("."):] if "." in name else ""
    if not data:
        raise ReadError("The file is empty.")
    if len(data) > MAX_BYTES:
        raise ReadError("The file is larger than 5 MB. Download a shorter date range and try again.")
    if ext in LATER:
        raise ReadError(f"{LATER[ext]} aren't supported yet. Download the statement as PDF, CSV or Excel instead.")
    if ext not in SUPPORTED:
        raise ReadError("Koshio can read PDF, CSV and Excel (.xlsx, .xls) statements. "
                        f"'{filename}' isn't one of those.")
    if ext == ".pdf":
        raise ReadError("Use read_statement for PDFs.")
    # Check the content really matches the extension rather than trusting the name.
    if ext == ".xlsx":
        if not data.startswith(b"PK"):
            raise ReadError("This file is named .xlsx but isn't an Excel workbook.")
        return _xlsx(data)
    if ext == ".xls":
        if data.startswith(b"\xd0\xcf\x11\xe0"):
            return _xls(data)
        if data.startswith(b"PK"):
            return _xlsx(data)
        # Many banks send HTML or text tables named .xls
        if b"<table" in data[:5000].lower() or b"<html" in data[:5000].lower():
            return _trim(_html_table(data))
        return _csv(data)
    if b"\x00" in data[:2000]:
        raise ReadError("This doesn't look like a text CSV file.")
    return _csv(data)


def _html_table(data: bytes) -> list[list[str]]:
    from html.parser import HTMLParser

    class P(HTMLParser):
        def __init__(self):
            super().__init__()
            self.rows, self.row, self.cell, self.in_cell = [], None, [], False

        def handle_starttag(self, tag, attrs):
            if tag == "tr":
                self.row = []
            elif tag in ("td", "th"):
                self.in_cell, self.cell = True, []

        def handle_endtag(self, tag):
            if tag in ("td", "th") and self.row is not None:
                self.row.append(" ".join("".join(self.cell).split()))
                self.in_cell = False
            elif tag == "tr" and self.row is not None:
                self.rows.append(self.row)
                self.row = None

        def handle_data(self, d):
            if self.in_cell:
                self.cell.append(d)

    p = P()
    p.feed(data.decode("utf-8", errors="replace"))
    if not p.rows:
        raise ReadError("No table was found in this file.")
    return p.rows


def read_statement(filename: str, data: bytes, password: "str | list[str] | None" = None,
                   with_password_index: bool = False):
    """Like read_rows, but also reads PDFs. Returns (rows, source suggested by the PDF reader),
    plus which of several candidate passwords worked when with_password_index is set."""
    name = (filename or "").lower()
    if name.endswith(".pdf") or data[:5] == b"%PDF-":
        if not data:
            raise ReadError("The file is empty.")
        if len(data) > MAX_PDF_BYTES:
            raise ReadError("The PDF is larger than 15 MB. Download a shorter date range and try again.")
        if not data.startswith(b"%PDF-"):
            raise ReadError("This file is named .pdf but isn't a PDF.")
        from .pdf import read_pdf
        rows, source, _style, used = read_pdf(data, password)
        return (_trim(rows), source, used) if with_password_index else (_trim(rows), source)
    rows = read_rows(filename, data)
    return (rows, None, None) if with_password_index else (rows, None)
